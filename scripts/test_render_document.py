"""Portable contract checks for both screen and formal document outputs."""

import copy
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from render_document import CSS, render_html, render_markdown, save_docx, validate  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / 'assets' / 'example-document.json'
CLI = ROOT / 'scripts' / 'render_document.py'


class DocumentDesignTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = validate(json.loads(SAMPLE.read_text(encoding='utf-8')))

    def test_screen_and_markdown_keep_main_content(self):
        webpage = render_html(self.data)
        markdown = render_markdown(self.data)
        for phrase in ('核心目标', 'Word / PDF', '读者与使用边界', '先减少无差别强调'):
            if phrase == '先减少无差别强调':
                self.assertIn(phrase, webpage)
                continue  # Markdown only presents choice labels, not private answer feedback.
            self.assertIn(phrase, webpage)
            self.assertIn(phrase, markdown)

    def test_inline_emphasis_and_html_escaping(self):
        webpage = render_html(self.data)
        markdown = render_markdown(self.data)
        self.assertIn('<strong>核心目标：</strong>', webpage)
        self.assertIn('<mark>屏幕阅读</mark>', webpage)
        self.assertIn('**屏幕阅读**', markdown)
        poisoned = copy.deepcopy(self.data)
        poisoned['title'] = '</title><script>alert(1)</script>'
        poisoned['sections'][0]['blocks'][0]['text'] = '**<img src=x onerror=alert(1)>**'
        escaped = render_html(validate(poisoned))
        self.assertNotIn('</title><script>alert', escaped)
        self.assertIn('&lt;img src=x', escaped)

    def test_html_semantics_and_real_choice(self):
        webpage = render_html(self.data)
        self.assertIn('<main><article>', webpage)
        self.assertEqual(webpage.count('<section id='), 4)
        self.assertIn('<th scope="col">', webpage)
        self.assertEqual(webpage.count('<input type="radio"'), 3)
        self.assertIn('type="submit" disabled', webpage)
        self.assertIn('role="status" aria-live="polite"', webpage)
        self.assertIn("feedback.textContent=selected.dataset.feedback", webpage)

    def test_visual_tokens_and_mobile(self):
        self.assertIn('--width:720px', CSS)
        self.assertIn('background:var(--surface)', CSS)
        self.assertIn('line-height:1.75', CSS)
        self.assertIn('@media(max-width:540px)', CSS)
        self.assertIn('@media print', CSS)

    def test_markdown_choice_is_honestly_static(self):
        markdown = render_markdown(self.data)
        self.assertIn('静态文档', markdown)
        self.assertNotIn('<button', markdown)
        self.assertIn('| 媒介 |', markdown)

    def test_rejects_bad_tables_and_duplicate_choice_ids(self):
        broken = copy.deepcopy(self.data)
        broken['sections'][2]['blocks'][0]['rows'][0].append('多余单元格')
        with self.assertRaises(ValueError):
            validate(broken)
        broken = copy.deepcopy(self.data)
        broken['sections'][3]['blocks'][1]['options'][1]['id'] = 'A'
        with self.assertRaises(ValueError):
            validate(broken)

    def test_rejects_invalid_schema_or_page_size(self):
        broken = copy.deepcopy(self.data)
        broken['schema_version'] = True
        with self.assertRaises(ValueError):
            validate(broken)
        broken = copy.deepcopy(self.data)
        broken['page_size'] = 'poster'
        with self.assertRaises(ValueError):
            validate(broken)

    def test_cli_md_html_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = subprocess.run([sys.executable, str(CLI), str(SAMPLE), '--out', tmp, '--formats', 'md,html'], capture_output=True, text=True, check=False)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue((Path(tmp) / 'example-document.md').is_file())
            self.assertTrue((Path(tmp) / 'example-document.html').is_file())
            duplicate = subprocess.run([sys.executable, str(CLI), str(SAMPLE), '--out', tmp, '--formats', 'md,html'], capture_output=True, text=True, check=False)
            self.assertEqual(duplicate.returncode, 2)
            self.assertIn('already exists', duplicate.stderr)

    def test_docx_title_styles_and_tables(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest('python-docx unavailable')
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / 'sample.docx'
            save_docx(self.data, dest)
            document = Document(str(dest))
            self.assertTrue(any(p.style.name == 'Title' and p.text == self.data['title'] for p in document.paragraphs))
            self.assertEqual(len([p for p in document.paragraphs if p.style.name == 'Heading 1']), 4)
            self.assertEqual(len(document.tables), 2)
            self.assertGreater(dest.stat().st_size, 5000)

    def test_pdf_cli_when_converter_is_available(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest('python-docx unavailable')
        if not (shutil.which('soffice') or shutil.which('libreoffice')):
            self.skipTest('LibreOffice unavailable')
        with tempfile.TemporaryDirectory() as tmp:
            result = subprocess.run([sys.executable, str(CLI), str(SAMPLE), '--out', tmp, '--formats', 'pdf'], capture_output=True, text=True, check=False, timeout=120)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue((Path(tmp) / 'example-document.pdf').stat().st_size > 2000)


if __name__ == '__main__':
    unittest.main()

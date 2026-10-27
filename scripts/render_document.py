#!/usr/bin/env python3
"""Render one semantic document to MD, HTML, DOCX, and optionally PDF.

HTML/Markdown: Python standard library. DOCX: python-docx. PDF: LibreOffice.
This deliberately small renderer does not replace an academic or legal typesetter.
"""

from __future__ import annotations

import argparse
import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

TYPES = {"paragraph", "facts", "steps", "list", "table", "note", "choice"}
FORMATS = {"md", "html", "docx", "pdf"}
INLINE = re.compile(r"\*\*(.+?)\*\*|==(.+?)==", re.DOTALL)

CSS = """
:root{color-scheme:light;--ink:#171717;--body:#303030;--muted:#686868;--surface:#f6f6f6;--line:#e8e8e8;--width:720px}
*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#fff;color:var(--body);font-family:-apple-system,BlinkMacSystemFont,'Noto Sans CJK SC','PingFang SC','Microsoft YaHei',sans-serif;font-size:16px;line-height:1.75;-webkit-font-smoothing:antialiased}
main{width:min(calc(100% - 48px),var(--width));margin:62px auto 100px}h1,h2,h3,p{overflow-wrap:anywhere}
h1{font-size:clamp(25px,3vw,29px);line-height:1.36;color:var(--ink);letter-spacing:-.01em;margin:0 0 9px;font-weight:680}
.subtitle{color:var(--muted);font-size:14px;margin:0 0 20px}.summary{font-size:19px;line-height:1.65;margin:0 0 25px;color:var(--ink)}
strong{color:var(--ink);font-weight:680}mark{background:#ececec;border-radius:4px;padding:0 .12em;color:inherit;box-decoration-break:clone;-webkit-box-decoration-break:clone}
.contents{font-size:14px;border-top:1px solid var(--line);border-bottom:1px solid var(--line);padding:12px 0;margin:20px 0 30px}.contents a{color:var(--body);text-decoration:none;margin-right:16px;display:inline-block}.contents a:hover{text-decoration:underline}
section{margin:34px 0}h2{font-size:20px;line-height:1.45;margin:0 0 14px;font-weight:680;color:var(--ink)}h3{font-size:17px;line-height:1.5;font-weight:660}
p{margin:0 0 14px;white-space:pre-wrap}.prose{line-height:1.78}.facts{background:var(--surface);border:1px solid var(--line);border-radius:14px;padding:4px 18px;margin:16px 0}
.fact{display:grid;grid-template-columns:minmax(85px,26%) minmax(0,1fr);gap:16px;padding:12px 0;border-bottom:1px solid #e3e3e3}.fact:last-child{border-bottom:0}.fact-label{color:var(--muted);font-size:13px}.fact-value{color:var(--ink)}
.steps{list-style:none;padding:0;margin:12px 0 16px}.step{display:grid;grid-template-columns:28px minmax(0,1fr);gap:12px;padding:0 0 16px}.step-num{display:flex;align-items:center;justify-content:center;width:26px;height:26px;border:1px solid var(--line);border-radius:50%;background:var(--surface);font-size:12px;color:var(--muted);font-weight:650}
.step-title{font-weight:650;color:var(--ink)}.step-detail{font-size:14px;color:var(--muted);margin-top:2px}.bullet-list{padding-left:22px;margin:12px 0 18px}.bullet-list li{margin:8px 0}
.table-wrap{overflow-x:auto;max-width:100%;margin:16px 0 20px}table{border-collapse:collapse;width:100%;font-size:14px;line-height:1.6}th,td{border:1px solid #dedede;padding:11px 12px;text-align:left;vertical-align:top;overflow-wrap:anywhere}th{background:var(--surface);font-weight:650;color:var(--ink)}
.note{border:1px solid var(--line);border-radius:14px;background:var(--surface);padding:14px 18px;margin:16px 0}.note-title{font-size:13px;color:var(--muted);margin:0 0 4px}.note-body{margin:0}
.quiz{border:1px solid var(--line);border-radius:14px;padding:18px;margin-top:16px}.quiz-question{font-weight:650;color:var(--ink);margin:0 0 10px}.quiz legend{font-size:13px;color:var(--muted);margin-bottom:8px}fieldset{border:0;padding:0;margin:0}.option{display:flex;align-items:flex-start;gap:12px;min-height:44px;padding:9px 8px;border:1px solid transparent;border-radius:9px;cursor:pointer}.option:hover,.option:has(input:checked){background:var(--surface)}.option input{width:19px;height:19px;accent-color:#171717;margin:4px 0 0;flex:none}.option span{flex:1}
button{display:block;width:100%;min-height:44px;border:0;border-radius:100px;background:#171717;color:white;font-family:inherit;font-size:15px;font-weight:650;cursor:pointer;margin-top:15px;padding:11px 16px}button:disabled{background:#ebebeb;color:#777;cursor:not-allowed}button:focus-visible,.option:focus-within,.contents a:focus-visible{outline:3px solid #256bb0;outline-offset:3px}.feedback{margin-top:15px;background:var(--surface);border-radius:10px;padding:12px 14px;white-space:pre-wrap}.feedback[hidden]{display:none}
@media(max-width:540px){main{width:min(calc(100% - 32px),var(--width));margin:32px auto 64px}.summary{font-size:18px}.facts{padding:4px 14px}.fact{grid-template-columns:86px minmax(0,1fr);gap:12px}.quiz{padding:14px}}
@page{size:letter;margin:24mm 25mm}@media print{body{font-size:11pt;line-height:1.5;-webkit-print-color-adjust:exact;print-color-adjust:exact}main{width:100%;margin:0}h1{font-size:22pt}h2{font-size:15pt;break-after:avoid}section{break-inside:auto}.facts,.note,.table-wrap{break-inside:avoid}.contents{display:none}button,.feedback{display:none}.quiz{break-inside:avoid}.summary{font-size:12pt}}
"""

JS = """
document.querySelectorAll('[data-choice-form]').forEach(form=>{
 const button=form.querySelector('button[type="submit"]');
 const feedback=form.querySelector('[data-feedback]');
 form.addEventListener('change',()=>{button.disabled=!form.querySelector('input:checked');feedback.textContent='';feedback.hidden=true;});
 form.addEventListener('submit',event=>{event.preventDefault();const selected=form.querySelector('input:checked');if(!selected)return;feedback.textContent=selected.dataset.feedback||'已提交';feedback.hidden=false;});
});
"""


def text_field(value: Any, field: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field} must be a nonempty string")
    return value.strip()


def bounded_list(value: Any, field: str, lo: int, hi: int) -> list[Any]:
    if not isinstance(value, list) or not lo <= len(value) <= hi:
        raise ValueError(f"{field} must contain {lo}-{hi} entries")
    return value


def validate(data: Any) -> dict[str, Any]:
    if not isinstance(data, dict) or type(data.get("schema_version")) is not int or data["schema_version"] != 1:
        raise ValueError("schema_version must equal integer 1")
    for key in ("title", "summary"):
        text_field(data.get(key), key)
    if data.get("subtitle") is not None:
        text_field(data["subtitle"], "subtitle")
    if not isinstance(data.get("page_size", "letter"), str) or data.get("page_size", "letter").lower() not in {"letter", "a4"}:
        raise ValueError("page_size must be letter or a4")
    sections = bounded_list(data.get("sections"), "sections", 1, 20)
    choice_count = 0
    for si, section in enumerate(sections):
        if not isinstance(section, dict):
            raise ValueError(f"sections[{si}] must be object")
        text_field(section.get("title"), f"sections[{si}].title")
        for bi, block in enumerate(bounded_list(section.get("blocks"), f"sections[{si}].blocks", 1, 18)):
            if not isinstance(block, dict) or block.get("type") not in TYPES:
                raise ValueError(f"sections[{si}].blocks[{bi}] has invalid type")
            kind = block["type"]
            ref = f"sections[{si}].blocks[{bi}]"
            if kind == "paragraph":
                text_field(block.get("text"), ref + ".text")
            elif kind == "note":
                text_field(block.get("title"), ref + ".title")
                text_field(block.get("body"), ref + ".body")
            elif kind in {"list", "steps", "facts"}:
                entries = bounded_list(block.get("items"), ref + ".items", 2, 8 if kind == "steps" else 5 if kind == "facts" else 20)
                for ii, entry in enumerate(entries):
                    if kind == "list":
                        text_field(entry, f"{ref}.items[{ii}]")
                    else:
                        if not isinstance(entry, dict):
                            raise ValueError(f"{ref}.items[{ii}] must be object")
                        for key in (("title", "detail") if kind == "steps" else ("label", "value")):
                            text_field(entry.get(key), f"{ref}.items[{ii}].{key}")
            elif kind == "table":
                headers = bounded_list(block.get("headers"), ref + ".headers", 2, 6)
                for hi, heading in enumerate(headers):
                    text_field(heading, f"{ref}.headers[{hi}]")
                rows = bounded_list(block.get("rows"), ref + ".rows", 1, 25)
                for ri, row in enumerate(rows):
                    if not isinstance(row, list) or len(row) != len(headers):
                        raise ValueError(f"{ref}.rows[{ri}] must match header column count")
                    for ci, cell in enumerate(row):
                        text_field(cell, f"{ref}.rows[{ri}][{ci}]")
            elif kind == "choice":
                choice_count += 1
                if choice_count > 1:
                    raise ValueError("only one choice block per document")
                text_field(block.get("question"), ref + ".question")
                seen: set[str] = set()
                for ii, option in enumerate(bounded_list(block.get("options"), ref + ".options", 2, 4)):
                    if not isinstance(option, dict):
                        raise ValueError(f"{ref}.options[{ii}] must be object")
                    option_id = text_field(option.get("id"), ref + ".option.id")
                    if not re.fullmatch(r"[A-Za-z0-9_-]{1,12}", option_id) or option_id in seen:
                        raise ValueError("invalid or duplicate option id")
                    seen.add(option_id)
                    text_field(option.get("label"), ref + ".option.label")
                    text_field(option.get("feedback"), ref + ".option.feedback")
    return data


def escape(value: str) -> str:
    return html.escape(value, quote=True)


def inline_html(value: str) -> str:
    result: list[str] = []
    end = 0
    for match in INLINE.finditer(value):
        result.append(escape(value[end:match.start()]))
        if match.group(1) is not None:
            result.append(f"<strong>{escape(match.group(1))}</strong>")
        else:
            result.append(f"<mark>{escape(match.group(2))}</mark>")
        end = match.end()
    result.append(escape(value[end:]))
    return "".join(result)


def inline_markdown(value: str) -> str:
    return re.sub(r"==(.+?)==", r"**\1**", value, flags=re.DOTALL)


def table_markdown(value: str) -> str:
    return inline_markdown(value).replace("|", r"\|").replace("\n", " ")


def render_block_html(block: dict[str, Any]) -> str:
    kind = block["type"]
    if kind == "paragraph":
        return f'<p class="prose">{inline_html(block["text"])}</p>'
    if kind == "facts":
        rows = "".join(f'<div class="fact"><span class="fact-label">{escape(x["label"])}</span><span class="fact-value">{inline_html(x["value"])}</span></div>' for x in block["items"])
        return f'<div class="facts">{rows}</div>'
    if kind == "steps":
        rows = "".join(f'<li class="step"><span class="step-num" aria-hidden="true">{i+1}</span><div><div class="step-title">{inline_html(x["title"])}</div><div class="step-detail">{inline_html(x["detail"])}</div></div></li>' for i, x in enumerate(block["items"]))
        return f'<ol class="steps">{rows}</ol>'
    if kind == "list":
        return '<ul class="bullet-list">' + "".join(f'<li>{inline_html(x)}</li>' for x in block["items"]) + '</ul>'
    if kind == "table":
        heads = "".join(f'<th scope="col">{inline_html(x)}</th>' for x in block["headers"])
        rows = "".join('<tr>' + "".join(f'<td>{inline_html(cell)}</td>' for cell in row) + '</tr>' for row in block["rows"])
        return f'<div class="table-wrap"><table><thead><tr>{heads}</tr></thead><tbody>{rows}</tbody></table></div>'
    if kind == "note":
        return f'<aside class="note"><p class="note-title">{escape(block["title"])}</p><p class="note-body">{inline_html(block["body"])}</p></aside>'
    options = "".join(f'<label class="option"><input type="radio" name="choice" value="{escape(x["id"])}" data-feedback="{escape(x["feedback"])}"><span>{escape(x["id"])}. {inline_html(x["label"])}</span></label>' for x in block["options"])
    return f'<div class="quiz"><p class="quiz-question">{inline_html(block["question"])}</p><form data-choice-form><fieldset><legend>请选择最符合你的选项</legend>{options}</fieldset><button type="submit" disabled>提交选择</button><div class="feedback" data-feedback role="status" aria-live="polite" hidden></div></form></div>'


def render_html(data: dict[str, Any]) -> str:
    contents = ""
    if len(data["sections"]) > 3:
        links = "".join(f'<a href="#s{i}">{escape(s["title"])}</a>' for i, s in enumerate(data["sections"]))
        contents = f'<nav class="contents" aria-label="目录">{links}</nav>'
    sections = "".join(f'<section id="s{i}"><h2>{escape(s["title"])}</h2>{"".join(render_block_html(b) for b in s["blocks"])}</section>' for i, s in enumerate(data["sections"]))
    subtitle = f'<p class="subtitle">{escape(data["subtitle"])}</p>' if data.get("subtitle") else ""
    css = CSS.replace('size:letter;', 'size:A4;' if data.get("page_size", "letter").lower() == "a4" else 'size:letter;')
    return ('<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>{escape(data["title"])}</title><style>{css}</style></head><body><main><article><h1>{escape(data["title"])}</h1>'
            f'{subtitle}<p class="summary">{inline_html(data["summary"])}</p>{contents}{sections}</article></main><script>{JS}</script></body></html>')


def render_markdown(data: dict[str, Any]) -> str:
    lines = [f'# {data["title"]}', ""]
    if data.get("subtitle"):
        lines.extend([data["subtitle"], ""])
    lines.extend([inline_markdown(data["summary"]), ""])
    for section in data["sections"]:
        lines.extend([f'## {section["title"]}', ""])
        for block in section["blocks"]:
            kind = block["type"]
            if kind == "paragraph":
                lines.append(inline_markdown(block["text"]))
            elif kind == "facts":
                lines += ['| 信息 | 内容 |', '| --- | --- |']
                lines += [f'| {table_markdown(x["label"])} | {table_markdown(x["value"])} |' for x in block["items"]]
            elif kind == "steps":
                lines += [f'{i+1}. **{inline_markdown(x["title"])}** — {inline_markdown(x["detail"])}' for i, x in enumerate(block["items"])]
            elif kind == "list":
                lines += [f'- {inline_markdown(x)}' for x in block["items"]]
            elif kind == "table":
                lines.append('| ' + ' | '.join(map(table_markdown, block['headers'])) + ' |')
                lines.append('| ' + ' | '.join('---' for _ in block['headers']) + ' |')
                lines += ['| ' + ' | '.join(map(table_markdown, row)) + ' |' for row in block['rows']]
            elif kind == "note":
                lines.append(f'> **{block["title"]}**  {inline_markdown(block["body"])}')
            else:
                lines.append(inline_markdown(block['question']))
                lines += [f'- {x["id"]}. {inline_markdown(x["label"])}' for x in block['options']]
                lines.append('（静态文档：请回复选项；此处没有可用的提交按钮。）')
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def add_docx_runs(paragraph: Any, value: str, *, always_bold: bool = False) -> None:
    position = 0
    for match in INLINE.finditer(value):
        if match.start() > position:
            fragment = paragraph.add_run(value[position:match.start()])
            fragment.bold = always_bold
        fragment = paragraph.add_run(match.group(1) if match.group(1) is not None else match.group(2))
        fragment.bold = True
        position = match.end()
    if position < len(value):
        fragment = paragraph.add_run(value[position:])
        fragment.bold = always_bold


def save_docx(data: dict[str, Any], dest: Path) -> None:
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Mm, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError('DOCX requires python-docx; install it or use --formats md,html') from exc

    document = Document()
    sec = document.sections[0]
    if data.get('page_size', 'letter').lower() == 'a4':
        sec.page_width, sec.page_height = Mm(210), Mm(297)
    else:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Inches(.95), Inches(.95), Inches(.97), Inches(.97)

    font_name = 'Noto Sans CJK SC'

    def set_style(name: str, size: float, bold: bool, before: float = 0, after: float = 6, line: float = 1.45) -> None:
        style = document.styles[name]
        style.font.name = font_name
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(23, 23, 23)
        r_pr = style._element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement('w:rFonts')
            r_pr.insert(0, r_fonts)
        for attribute in ('ascii', 'hAnsi', 'eastAsia'):
            r_fonts.set(qn('w:' + attribute), font_name)
        pf = style.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = line

    set_style('Normal', 11, False, after=6)
    set_style('Title', 23, True, after=9, line=1.3)
    set_style('Subtitle', 10, False, after=16)
    set_style('Heading 1', 15.5, True, before=16, after=7, line=1.3)
    set_style('Heading 2', 12.5, True, before=10, after=5)
    set_style('List Bullet', 11, False, after=5)
    set_style('List Number', 11, False, after=5)
    # Some Word templates put a blue bottom rule on the built-in Title style.
    # The visual contract asks for an unboxed black title with whitespace only.
    title_properties = document.styles['Title']._element.get_or_add_pPr()
    for old_border in list(title_properties.findall(qn('w:pBdr'))):
        title_properties.remove(old_border)
    for style_name in ('Title', 'Heading 1', 'Heading 2'):
        document.styles[style_name].paragraph_format.keep_with_next = True
    document.styles['Normal'].paragraph_format.widow_control = True

    title = document.add_paragraph(style='Title')
    title.add_run(data['title'])
    if data.get('subtitle'):
        document.add_paragraph(data['subtitle'], style='Subtitle')
    summary = document.add_paragraph()
    summary.paragraph_format.space_after = Pt(15)
    add_docx_runs(summary, data['summary'])

    def disable_row_split(row: Any) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement('w:cantSplit'))

    def shade(cell: Any) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'F5F5F5')
        tc_pr.append(shd)

    def set_table_borders(table: Any) -> None:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement('w:' + edge)
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), '4')
            el.set(qn('w:color'), 'D9D9D9')
            borders.append(el)
        tbl_pr.append(borders)

    for section in data['sections']:
        document.add_paragraph(section['title'], style='Heading 1')
        for block in section['blocks']:
            kind = block['type']
            if kind == 'paragraph':
                add_docx_runs(document.add_paragraph(), block['text'])
            elif kind == 'note':
                p = document.add_paragraph()
                p.add_run(block['title'] + '  ').bold = True
                add_docx_runs(p, block['body'])
            elif kind == 'list':
                for item in block['items']:
                    add_docx_runs(document.add_paragraph(style='List Bullet'), item)
            elif kind == 'steps':
                for item in block['items']:
                    p = document.add_paragraph(style='List Number')
                    add_docx_runs(p, item['title'], always_bold=True)
                    add_docx_runs(p, ' — ' + item['detail'])
            elif kind in {'facts', 'table'}:
                if kind == 'facts':
                    rows = [[x['label'], x['value']] for x in block['items']]
                    headers = None
                else:
                    headers, rows = block['headers'], block['rows']
                table = document.add_table(rows=(1 if headers else 0) + len(rows), cols=len(headers) if headers else 2)
                table.autofit = False
                set_table_borders(table)
                if headers:
                    for col, celltext in enumerate(headers):
                        cell = table.rows[0].cells[col]
                        shade(cell)
                        add_docx_runs(cell.paragraphs[0], celltext, always_bold=True)
                    tr_pr = table.rows[0]._tr.get_or_add_trPr()
                    flag = OxmlElement('w:tblHeader')
                    flag.set(qn('w:val'), 'true')
                    tr_pr.append(flag)
                for ri, values in enumerate(rows):
                    row = table.rows[ri + (1 if headers else 0)]
                    disable_row_split(row)
                    for ci, value in enumerate(values):
                        cell = row.cells[ci]
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        if kind == 'facts' and ci == 0:
                            shade(cell)
                        add_docx_runs(cell.paragraphs[0], value, always_bold=(kind == 'facts' and ci == 0))
                if kind == 'facts':
                    available = (8.5 - 1.94) if data.get('page_size', 'letter').lower() == 'letter' else (210 / 25.4 - 1.94)
                    for row in table.rows:
                        row.cells[0].width = Inches(available * .27)
                        row.cells[1].width = Inches(available * .73)
                document.add_paragraph().paragraph_format.space_after = Pt(2)
            else:  # choices are intentionally a static questionnaire in Word
                p = document.add_paragraph()
                add_docx_runs(p, block['question'], always_bold=True)
                for option in block['options']:
                    document.add_paragraph(f'({option["id"]}) {inline_markdown(option["label"]).replace("**", "")}', style='List Bullet')
                note = document.add_paragraph('静态文档：请标记选项；此处没有网页提交功能。')
                note.paragraph_format.space_after = Pt(10)
    document.save(str(dest))


def export_pdf(docx_path: Path, pdf_path: Path) -> None:
    binary = shutil.which('soffice') or shutil.which('libreoffice')
    if not binary:
        raise RuntimeError('PDF export requires LibreOffice/soffice; generate DOCX and export PDF manually')
    with tempfile.TemporaryDirectory(prefix='document-design-lo-') as profile:
        profile_uri = Path(profile).resolve().as_uri()
        cmd = [binary, f'-env:UserInstallation={profile_uri}', '--headless', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(pdf_path.parent), str(docx_path)]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=100, check=False)
        expected = pdf_path.parent / (docx_path.stem + '.pdf')
        if result.returncode != 0 or not expected.is_file():
            raise RuntimeError(f'PDF conversion failed: {(result.stderr or result.stdout).strip()}')
        if expected.resolve() != pdf_path.resolve():
            expected.replace(pdf_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--formats', default='md,html', help='Comma-separated: md,html,docx,pdf')
    parser.add_argument('--force', action='store_true', help='Overwrite existing generated outputs')
    args = parser.parse_args()
    formats = [x.strip() for x in args.formats.split(',')]
    if not formats or len(set(formats)) != len(formats) or set(formats) - FORMATS:
        parser.error('Invalid or duplicate --formats entries')
    try:
        data = validate(json.loads(args.source.read_text(encoding='utf-8')))
        out = args.out.resolve()
        stem = args.source.stem
        targets = {fmt: out / f'{stem}.{fmt}' for fmt in formats}
        for path in targets.values():
            if path.exists() and not args.force:
                raise FileExistsError(f'{path} already exists (use --force)')
        out.mkdir(parents=True, exist_ok=True)
        if 'md' in targets:
            targets['md'].write_text(render_markdown(data), encoding='utf-8')
        if 'html' in targets:
            targets['html'].write_text(render_html(data), encoding='utf-8')
        if 'docx' in targets or 'pdf' in targets:
            if 'docx' in targets:
                docx_path = targets['docx']
                save_docx(data, docx_path)
                if 'pdf' in targets:
                    export_pdf(docx_path, targets['pdf'])
            else:
                with tempfile.TemporaryDirectory(prefix='document-design-pdf-') as td:
                    docx_path = Path(td) / f'{stem}.docx'
                    save_docx(data, docx_path)
                    export_pdf(docx_path, targets['pdf'])
    except (OSError, ValueError, RuntimeError, json.JSONDecodeError, subprocess.TimeoutExpired) as exc:
        print(f'render error: {exc}', file=sys.stderr)
        return 2
    for fmt in formats:
        print(f'{fmt.upper()}: {targets[fmt]}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
